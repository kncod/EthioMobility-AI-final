"""Build the 5-slide hackathon presentation (Deliverable F)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
OUT = ROOT / "presentation" / "team_addis_demand_ai_slides.pptx"
TEAM = "Addis Demand AI"
MEMBERS = "Abraham Getachew · Ahmed Hussen · Natanim Masresha · Nigus Shiferaw · Tekilu Asefa"
TRACK = "IDAE · EDI-Qiyas-CoDiST Advanced Digital Skills Program"

# Widescreen 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x0B, 0x3D, 0x5C)
ACCENT = RGBColor(0x00, 0x72, 0xB2)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)


def set_run(run, text, size=18, bold=False, color=DARK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_banner(slide, title: str):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.05))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    set_run(p.add_run(), title, size=26, bold=True, color=WHITE)
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.25)


def add_textbox(slide, left, top, width, height, lines, size=16, color=DARK, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.clear()
        set_run(p.add_run(), line, size=size, bold=bold, color=color)
        p.space_after = Pt(6)
    return box


def add_bullets(slide, left, top, width, height, items, size=17):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.clear()
        set_run(p.add_run(), f"•  {item}", size=size, color=DARK)
        p.space_after = Pt(8)
        p.level = 0
    return box


def set_notes(slide, text: str):
    notes = slide.notes_slide.notes_text_frame
    notes.text = text


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ----- Slide 1 -----
    s1 = prs.slides.add_slide(blank)
    add_banner(s1, f"1 · {TEAM} — Addis ride demand forecasting")
    add_textbox(
        s1,
        Inches(0.5),
        Inches(1.2),
        Inches(12.2),
        Inches(0.9),
        [MEMBERS, TRACK],
        size=14,
        color=MUTED,
    )
    add_bullets(
        s1,
        Inches(0.5),
        Inches(2.1),
        Inches(12.2),
        Inches(4.8),
        [
            "Goal: forecast hourly trips for 12 Addis Ababa zones, 1–14 Nov 2025",
            "History: 1 Jan – 31 Oct 2025 (~85k zone-hours)",
            "Three messy tables to integrate — not cleaned for us:",
            "     ① ride_demand (trips by zone × hour)   ② weather_hourly (city)   ③ events_calendar (intervals)",
            "Why it matters: tell drivers where to be, a day or two ahead",
            "Key constraint: validate by time; only use features known at forecast time",
        ],
        size=18,
    )
    set_notes(
        s1,
        "Open with the ops problem. Emphasize three raw tables and that joining clocks/zones is half the work. "
        "Mention synthetic data for training.",
    )

    # ----- Slide 2 -----
    s2 = prs.slides.add_slide(blank)
    add_banner(s2, "2 · Cleaning & integration — clocks and joins")
    add_bullets(
        s2,
        Inches(0.45),
        Inches(1.3),
        Inches(6.4),
        Inches(5.8),
        [
            "55 zone spellings → 12 canonical labels",
            "Clock proof: trips are EAT (+03); weather is mostly UTC (Z)",
            "Convert weather UTC→Africa/Addis_Ababa before join",
            "Weather join: many-to-one on hour — match rate 100% after hourly spine",
            "Events: interval join; football/concert ±2h; confirmed only (159 matched)",
            "Sentinels fixed (rain −9999); negatives/caps logged in A1",
        ],
        size=17,
    )
    fig06 = FIG / "fig06_weather_timezone_check.png"
    if fig06.exists():
        s2.shapes.add_picture(str(fig06), Inches(7.0), Inches(1.35), width=Inches(5.9))
    set_notes(
        s2,
        "Walk the join map left-to-right. Stress the +3h clock bug — wrong clock kills the rain signal. "
        "Point at fig06: afternoon temp peak only after conversion.",
    )

    # ----- Slide 3 -----
    s3 = prs.slides.add_slide(blank)
    add_banner(s3, "3 · What the data says — two findings")
    add_bullets(
        s3,
        Inches(0.4),
        Inches(1.25),
        Inches(5.8),
        Inches(2.2),
        [
            "① Demand grew ~30% Jan→Oct — November needs trend/lags, not only seasonality",
            "② Football: largest uplift in the 2 hours before kickoff (~2.7× baseline in stadium zones)",
        ],
        size=16,
    )
    fig03 = FIG / "fig03_demand_trend_with_holidays.png"
    fig08 = FIG / "fig08_event_study.png"
    if fig03.exists():
        s3.shapes.add_picture(str(fig03), Inches(0.35), Inches(3.5), width=Inches(6.2))
    if fig08.exists():
        s3.shapes.add_picture(str(fig08), Inches(6.7), Inches(3.5), width=Inches(6.2))
    add_textbox(
        s3,
        Inches(6.5),
        Inches(1.3),
        Inches(6.4),
        Inches(2.0),
        [
            "Also: Ayat is a late-launch zone (shorter history).",
            "Rain lifts demand unevenly by zone type.",
            "Cancelled events excluded from features.",
        ],
        size=14,
        color=MUTED,
    )
    set_notes(
        s3,
        "Two takeaways only. Growth → week_index + lags. Football before-window justifies ±2h features. "
        "These are from Deliverable B/C figures.",
    )

    # ----- Slide 4 -----
    s4 = prs.slides.add_slide(blank)
    add_banner(s4, "4 · Modeling & evaluation — beat the seasonal baseline")
    add_bullets(
        s4,
        Inches(0.45),
        Inches(1.25),
        Inches(6.3),
        Inches(5.8),
        [
            "Split: train < 18 Oct · validate 18–31 Oct (time-ordered)",
            "Baselines: mean RMSE 27.9 · seasonal-naive 12.1",
            "Winner: HistGBM tuned → RMSE 10.16 · MAE 6.41",
            "Rolling origin (5 folds): 11.06 ± 0.42 RMSE",
            "Ablation (tuned): calendar 10.43 → +events+weather 10.16 (−0.27)",
            "No leakage: fare / wait / drivers excluded from inputs",
        ],
        size=17,
    )
    fig10 = FIG / "fig10_model_comparison.png"
    if fig10.exists():
        s4.shapes.add_picture(str(fig10), Inches(6.9), Inches(1.35), width=Inches(6.0))
    set_notes(
        s4,
        "Seasonal naive is the bar. HistGBM clears it. Ablation shows joins were worth it with tuned params. "
        "Mention Rule 5/6/7 briefly.",
    )

    # ----- Slide 5 -----
    s5 = prs.slides.add_slide(blank)
    add_banner(s5, "5 · Errors, score, demo — and what’s next")
    add_bullets(
        s5,
        Inches(0.45),
        Inches(1.3),
        Inches(12.2),
        Inches(5.8),
        [
            "Errors concentrate on peak hours, holidays, and event spikes (some unlisted in the calendar)",
            "Tried peak/rain interactions after D7 — did not help; we kept the simpler tuned model",
            "Final validation: RMSE ≈ 10.2 trips/zone-hour (~MAE 6.4) → ~±5 drivers buffer at 1.3 trips/driver-hour",
            "Demo (local):  python -m streamlit run app/app.py   →  http://localhost:8501",
            f"Submission: 4,032 rows · team_addis_demand_ai_submission.csv · Team: {TEAM}",
            "Next: quantile intervals for ops buffers · zone-specific models for Ayat · monitor drift after launch",
        ],
        size=18,
    )
    set_notes(
        s5,
        "Live demo: pick Kazanchis on a match day (e.g. 9 Nov) and show weather/event lookup + drivers. "
        "Close with honest next steps. Invite judges to choose zone/date.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
